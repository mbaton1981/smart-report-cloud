import hashlib
import io
import os
import sqlite3
import subprocess
import sys
import tempfile
import unittest
import uuid
from contextlib import closing
from datetime import date
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

from flask import Flask
from openpyxl import load_workbook
from pypdf import PdfReader

from object_reports import (
    ACCOUNTING_NOTE_RU, TYPE_LABELS, build_report, register_object_reports, report_objects, report_period,
)


ROOT = Path(__file__).resolve().parents[1]


class ObjectReportTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(prefix="smartreport-object-report-")
        self.addCleanup(self.directory.cleanup)
        self.database = Path(self.directory.name) / "server-copy.db"
        with closing(sqlite3.connect(self.database)) as conn:
            conn.executescript("""
                CREATE TABLE meta_objects(
                    id INTEGER PRIMARY KEY, name TEXT, company TEXT, markning TEXT, is_active INTEGER
                );
                CREATE TABLE incoming_submissions(
                    id INTEGER PRIMARY KEY, sync_id TEXT UNIQUE, date TEXT, employee TEXT, company TEXT,
                    object_name TEXT, submission_type TEXT, hours REAL, is_deleted INTEGER DEFAULT 0,
                    rate REAL DEFAULT 888888, comment TEXT DEFAULT 'PRIVATE-COMMENT-DO-NOT-EXPORT',
                    start_time TEXT, end_time TEXT
                );
                INSERT INTO meta_objects VALUES (1, 'Testvägen 10', 'Testkund AB', 'T-001', 1);
                INSERT INTO meta_objects VALUES (2, 'Testvägen 10', 'Annan kund AB', 'T-002', 1);
                INSERT INTO meta_objects VALUES (3, 'Semestertimmar', 'Vårt AB', 'p1002', 1);
                INSERT INTO meta_objects VALUES (4, 'Inaktivt objekt', 'Testkund AB', 'T-003', 0);
                CREATE TABLE meta_employees(name TEXT, is_active INTEGER);
                INSERT INTO meta_employees VALUES ('Erik Öberg', 0);
                CREATE TABLE cloud_shifts(date TEXT, employee TEXT, object_name TEXT, hours REAL);
                INSERT INTO cloud_shifts VALUES ('2026-10-01', 'Anna Åström', 'Testvägen 10', 99999);
            """)
            conn.commit()
        self.app = Flask(__name__, template_folder=str(ROOT / "templates"))
        self.app.config.update(TESTING=True, SECRET_KEY="isolated-report-test-only")
        register_object_reports(self.app, lambda: str(self.database))
        self.client = self.app.test_client()
        self.login()
        with closing(self.connection()) as conn:
            self.project = next(p for p in report_objects(conn) if p.marking == "T-001")
        self.add_shift("2026-10-01", "Anna Åström", "work", 7.5)
        self.add_shift("2026-10-01", " Anna  Åström ", "work", 0.5)
        self.add_shift("2026-10-02", "Anna Åström", "ob", 2)
        self.add_shift("2026-10-03", "Erik Öberg", "work", 6.25)
        self.add_shift("2026-09-30", "Anna Åström", "work", 24)
        self.add_shift("2026-11-01", "Anna Åström", "work", 24)
        self.add_shift("2026-10-01", "Anna Åström", "work", 23, deleted=1)
        self.add_shift("2026-10-01", "Anna Åström", "work", 22, company="Annan kund AB")
        self.add_shift("2026-10-01", "Anna Åström", "vacation", 8, company="Vårt AB", project="Semestertimmar")

    def connection(self):
        conn = sqlite3.connect(self.database)
        conn.row_factory = sqlite3.Row
        return conn

    def login(self, role="admin"):
        with self.client.session_transaction() as session:
            session.update(user_id=1, role=role)

    def add_shift(self, day, employee, kind, hours, *, company="Testkund AB", project="Testvägen 10", deleted=0,
                  start_time=None, end_time=None):
        with closing(self.connection()) as conn:
            conn.execute("""
                INSERT INTO incoming_submissions
                    (sync_id,date,employee,company,object_name,submission_type,hours,is_deleted,start_time,end_time)
                VALUES (?,?,?,?,?,?,?,?,?,?)
            """, (str(uuid.uuid4()), day, employee, company, project, kind, hours, deleted,
                  start_time or ("18:00" if kind == "ob" else None),
                  end_time or ("20:00" if kind == "ob" else None)))
            conn.commit()

    def params(self, **extra):
        return {"project": self.project.key, "preset": "custom", "start": "2026-10-01", "end": "2026-10-31", **extra}

    def fetch_report(self, project=None):
        with closing(self.connection()) as conn:
            return build_report(conn, project or self.project, date(2026, 10, 1), date(2026, 10, 31))

    def test_aggregates_dates_staff_and_types_without_deleted_other_projects_or_mirror(self):
        report = self.fetch_report()
        self.assertEqual(len(report.rows), 3)
        self.assertEqual([e.employee for e in report.employees], ["Anna Åström", "Erik Öberg"])
        self.assertEqual(report.rows[0].hours, Decimal("8.00"))
        self.assertEqual(report.employees[0].total, Decimal("10.00"))
        self.assertEqual(report.totals.hours["work"], Decimal("14.25"))
        self.assertEqual(report.totals.hours["ob"], Decimal("2.00"))
        self.assertEqual(report.totals.hours["vacation"], Decimal("0.00"))
        self.assertEqual(report.totals.total, Decimal("16.25"))

    def test_all_six_types_are_separate_and_total_is_sum_of_records(self):
        for kind in ("atf", "sick", "vacation", "holiday"):
            self.add_shift("2026-10-05", "Anna Åström", kind, 1.5)
        report = self.fetch_report()
        self.assertEqual(report.totals.total, Decimal("22.25"))
        for kind in ("atf", "sick", "vacation", "holiday"):
            self.assertEqual(report.totals.hours[kind], Decimal("1.50"))
        response = self.client.get("/object-report.xlsx", query_string=self.params())
        summary = load_workbook(io.BytesIO(response.data))["Sammanfattning"]
        self.assertEqual([summary.cell(summary.max_row, col).value for col in range(2, 9)], [
            14.25, 2, 1.5, 1.5, 1.5, 1.5, 22.25,
        ])
        response = self.client.get("/object-report.pdf", query_string=self.params())
        text = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(response.data)).pages)
        self.assertIn("Totalt registrerade timmar: 22.25 h", text)
        self.assertNotIn("Faktiskt arbete", text)
        for label in TYPE_LABELS.values():
            self.assertIn(label, text)

    def test_period_presets_and_inclusive_custom_dates(self):
        for preset, anchor, end, expected in [
            ("week", "2026-10-20", "", ("2026-10-19", "2026-10-25")),
            ("fortnight", "2026-12-31", "", ("2026-12-28", "2027-01-10")),
            ("month", "2028-02-20", "", ("2028-02-01", "2028-02-29")),
            ("custom", "2026-10-01", "2026-10-01", ("2026-10-01", "2026-10-01")),
        ]:
            with self.subTest(preset=preset):
                self.assertEqual(tuple(d.isoformat() for d in report_period(preset, anchor, end)), expected)
        response = self.client.get("/object-report", query_string=self.params(end="2026-10-01"))
        self.assertEqual(response.status_code, 200)
        self.assertIn('id="all-total">8.00', response.get_data(as_text=True))

    def test_invalid_dates_and_object_return_explicit_errors_in_all_formats(self):
        for extra in [
            {"start": "2026-02-30"}, {"end": "2026-09-01"}, {"preset": "bad"},
            {"start": "20261001"}, {"project": "unknown"}, {"end": ""},
            {"start": "9999-12-31", "preset": "fortnight"},
        ]:
            for suffix in ("", ".pdf", ".xlsx"):
                response = self.client.get("/object-report" + suffix, query_string=self.params(**extra))
                self.assertEqual(response.status_code, 400, (extra, suffix))
        for suffix in (".pdf", ".xlsx"):
            self.assertEqual(self.client.get("/object-report" + suffix).status_code, 400)

    def test_report_requires_admin_for_view_and_all_exports(self):
        for role, status in (("user", 403), ("admin", 200)):
            self.login(role)
            for suffix in ("", ".pdf", ".xlsx"):
                self.assertEqual(self.client.get(
                    "/object-report" + suffix, query_string=self.params()
                ).status_code, status)
        with self.client.session_transaction() as session:
            session.clear()
        for suffix in ("", ".pdf", ".xlsx"):
            response = self.client.get("/object-report" + suffix, query_string=self.params())
            self.assertEqual(response.status_code, 401)
            self.assertNotIn("Testvägen", response.get_data(as_text=True))

    def test_old_names_are_visible_without_guessing_links_or_merging_another_company(self):
        with closing(self.connection()) as conn:
            conn.execute("UPDATE meta_objects SET name='Nytt namn' WHERE id=1")
            conn.commit()
            projects = report_objects(conn)
            historic = next(p for p in projects if p.name == "Testvägen 10" and p.company == "Testkund AB")
            current = next(p for p in projects if p.name == "Nytt namn")
        self.assertEqual(historic.state, "только история")
        self.assertEqual(historic.key, self.project.key)
        self.assertEqual(self.fetch_report(historic).totals.total, Decimal("16.25"))
        self.assertEqual(self.fetch_report(current).totals.total, Decimal("0"))
        response = self.client.get("/object-report")
        self.assertIn("не объединяются автоматически", response.get_data(as_text=True))

    def test_inactive_objects_and_inactive_employees_remain_reportable(self):
        self.add_shift("2026-10-02", "Erik Öberg", "work", 4, project="Inaktivt objekt")
        with closing(self.connection()) as conn:
            project = next(p for p in report_objects(conn) if p.marking == "T-003")
        self.assertEqual(project.state, "неактивен")
        self.assertEqual(self.fetch_report(project).totals.total, Decimal("4.00"))

    def test_html_pdf_and_xlsx_have_same_hours_and_no_private_fields(self):
        html = self.client.get("/object-report", query_string=self.params())
        self.assertEqual(html.status_code, 200)
        text = html.get_data(as_text=True)
        self.assertIn('id="all-total">16.25', text)
        self.assertIn("Всего учтено часов по записям", text)
        self.assertIn(ACCOUNTING_NOTE_RU, text)
        self.assertNotIn("Фактически отработано", text)
        self.assertIn("Неотправленные", text)
        self.assertIn("Скачать Excel", text)
        pdf = self.client.get("/object-report.pdf", query_string=self.params())
        self.assertEqual(pdf.status_code, 200)
        self.assertEqual(pdf.mimetype, "application/pdf")
        self.assertIn("attachment;", pdf.headers["Content-Disposition"])
        pdf_text = "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(pdf.data)).pages)
        self.assertIn("Totalt registrerade timmar: 16.25 h", pdf_text)
        self.assertIn("Anna Åström", pdf_text)
        self.assertIn("Testvägen 10", pdf_text)
        self.assertIn("2026-10-01", pdf_text)
        self.assertIn("ännu inte skickats", pdf_text)
        xlsx = self.client.get("/object-report.xlsx", query_string=self.params())
        self.assertEqual(xlsx.status_code, 200)
        workbook = load_workbook(io.BytesIO(xlsx.data))
        summary, detail = workbook.worksheets
        self.assertEqual(summary.cell(summary.max_row, 8).value, 16.25)
        self.assertEqual(summary.max_column, 8)
        self.assertEqual(sum(detail.cell(row, 4).value for row in range(2, detail.max_row + 1)), 16.25)
        self.assertTrue(all(detail.cell(row, 4).data_type == "n" for row in range(2, detail.max_row + 1)))
        self.assertIn("ännu inte skickats", summary["A5"].value)
        xlsx_text = str([[cell.value for row in sheet for cell in row] for sheet in workbook])
        for content in (text, pdf_text, xlsx_text):
            self.assertNotIn("PRIVATE-COMMENT-DO-NOT-EXPORT", content)
            self.assertNotIn("888888", content)
        for response in (html, pdf, xlsx):
            self.assertEqual(response.headers["Cache-Control"], "no-store, private")

    def test_displayed_rounding_and_totals_agree(self):
        for hours in (0.005, 0.005, 0.004):
            self.add_shift("2026-10-10", "Anna Åström", "work", hours)
        self.add_shift("2026-10-11", "Anna Åström", "work", 0.005)
        report = self.fetch_report()
        self.assertEqual(report.totals.total, Decimal("16.27"))
        self.assertEqual(sum((r.hours for r in report.rows), Decimal(0)), report.totals.total)

    def test_cross_project_ob_warning_in_all_formats_without_disclosing_other_project(self):
        self.add_shift("2026-10-02", " anna  åSTRÖM ", "ob", 2,
                       company="PRIVATE-OTHER-CUSTOMER", project="PRIVATE-OTHER-PROJECT",
                       start_time="19:00", end_time="21:00")
        with closing(self.connection()) as conn:
            before = "\n".join(conn.iterdump())
        report = self.fetch_report()
        self.assertTrue(report.ob_overlap)
        self.assertFalse(report.ob_check_incomplete)
        self.assertEqual(report.totals.total, Decimal("16.25"))
        self.assertEqual(report.totals.hours["ob"], Decimal("2.00"))
        html = self.client.get("/object-report", query_string=self.params()).get_data(as_text=True)
        html_report = html.split('id="report"', 1)[1]
        self.assertIn("обнаружены пересекающиеся OB-интервалы", html_report)
        self.assertIn("Проверка выполнена по всем объектам", html_report)
        pdf = self.client.get("/object-report.pdf", query_string=self.params())
        pdf_text = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf.data)).pages)
        xlsx = self.client.get("/object-report.xlsx", query_string=self.params())
        book = load_workbook(io.BytesIO(xlsx.data))
        xlsx_text = str([list(s.values) for s in book])
        for content in (pdf_text, xlsx_text):
            self.assertIn("överlappande OB-intervall", content)
            self.assertIn("vanligt arbete saknar start- och sluttid", content)
        for content in (html_report, pdf_text, xlsx_text):
            self.assertNotIn("PRIVATE-OTHER", content)
            self.assertNotIn("19:00", content)
            self.assertNotIn("21:00", content)
        self.assertEqual(book["Sammanfattning"].cell(book["Sammanfattning"].max_row, 8).value, 16.25)
        with closing(self.connection()) as conn:
            self.assertEqual(before, "\n".join(conn.iterdump()))

    def test_same_project_overlap_is_warned_without_combining_or_deducting_hours(self):
        self.add_shift("2026-10-02", "Anna Åström", "ob", 2, start_time="19:00", end_time="21:00")
        report = self.fetch_report()
        self.assertTrue(report.ob_overlap)
        self.assertEqual(report.totals.hours["ob"], Decimal("4.00"))
        self.assertEqual(report.totals.total, Decimal("18.25"))

    def test_touching_deleted_different_employee_and_unrelated_conflicts_do_not_warn(self):
        for employee, begin, finish, deleted in [
            ("Anna Åström", "20:00", "22:00", 0),
            ("Anna Åström", "17:00", "18:00", 0),
            ("Erik Öberg", "19:00", "21:00", 0),
            ("Anna Åström", "19:00", "21:00", 1),
        ]:
            self.add_shift("2026-10-02", employee, "ob", 2, company="Other",
                           project="Other", start_time=begin, end_time=finish, deleted=deleted)
        for begin, finish in [("08:00", "12:00"), ("10:00", "11:00")]:
            self.add_shift("2026-10-02", "Anna Åström", "ob", 1, company="Other",
                           project="Other", start_time=begin, end_time=finish)
        self.assertFalse(self.fetch_report().ob_overlap)
        self.assertFalse(self.fetch_report().ob_check_incomplete)

    def test_overnight_ob_checks_previous_and_next_day_across_period_boundaries(self):
        self.add_shift("2026-10-01", "Anna Åström", "ob", 1, start_time="00:00", end_time="01:00")
        self.add_shift("2026-09-30", "Anna Åström", "ob", 2, company="Other", project="Other",
                       start_time="23:00", end_time="01:00")
        self.assertTrue(self.fetch_report().ob_overlap)
        with closing(self.connection()) as conn:
            conn.execute("DELETE FROM incoming_submissions WHERE submission_type='ob'")
            conn.commit()
        self.add_shift("2026-10-31", "Anna Åström", "ob", 2, start_time="23:00", end_time="01:00")
        self.add_shift("2026-11-01", "Anna Åström", "ob", 2, company="Other", project="Other",
                       start_time="00:30", end_time="02:30")
        self.assertTrue(self.fetch_report().ob_overlap)

    def test_contained_and_identical_intervals_warn_but_midnight_touch_does_not(self):
        for begin, finish in [("18:30", "19:00"), ("18:00", "20:00")]:
            with self.subTest(begin=begin):
                self.add_shift("2026-10-02", "Anna Åström", "ob", 0.5, company="Other",
                               project="Other", start_time=begin, end_time=finish)
                self.assertTrue(self.fetch_report().ob_overlap)
                with closing(self.connection()) as conn:
                    conn.execute("DELETE FROM incoming_submissions WHERE company='Other'")
                    conn.commit()
        self.add_shift("2026-10-03", "Anna Åström", "ob", 2, start_time="22:00", end_time="24:00")
        self.add_shift("2026-10-04", "Anna Åström", "ob", 2, company="Other", project="Other",
                       start_time="00:00", end_time="02:00")
        self.assertFalse(self.fetch_report().ob_overlap)

    def test_missing_or_invalid_ob_times_produce_incomplete_warning_not_false_assurance(self):
        for begin, finish in [(None, None), ("25:00", "26:00"), ("18:00", "18:00")]:
            with closing(self.connection()) as conn:
                conn.execute("UPDATE incoming_submissions SET start_time=?,end_time=? WHERE submission_type='ob'",
                             (begin, finish))
                conn.commit()
            report = self.fetch_report()
            self.assertTrue(report.ob_check_incomplete)
            self.assertEqual(report.totals.total, Decimal("16.25"))
            html = self.client.get("/object-report", query_string=self.params()).get_data(as_text=True)
            self.assertIn("проверка пересечений OB неполная", html)
            pdf = self.client.get("/object-report.pdf", query_string=self.params())
            self.assertIn("ofullständig", "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf.data)).pages))
            xlsx = self.client.get("/object-report.xlsx", query_string=self.params())
            self.assertIn("ofullständig", str(list(load_workbook(io.BytesIO(xlsx.data))["Sammanfattning"].values)))

    def test_read_only_connection_and_database_unchanged_by_all_outputs(self):
        before = hashlib.sha256(self.database.read_bytes()).hexdigest()
        traces = []
        original = sqlite3.connect

        def connect(*args, **kwargs):
            self.assertTrue(kwargs.get("uri"))
            self.assertTrue(args[0].endswith("?mode=ro"))
            conn = original(*args, **kwargs)
            conn.set_trace_callback(traces.append)
            return conn

        with patch("object_reports.sqlite3.connect", side_effect=connect):
            for suffix in ("", ".pdf", ".xlsx"):
                response = self.client.get("/object-report" + suffix, query_string=self.params())
                self.assertEqual(response.status_code, 200)
        self.assertTrue(any("PRAGMA query_only=ON" in trace for trace in traces))
        self.assertFalse(any(trace.startswith(("INSERT", "UPDATE", "DELETE", "CREATE")) for trace in traces))
        self.assertEqual(hashlib.sha256(self.database.read_bytes()).hexdigest(), before)

    def test_excel_names_are_literal_not_formulas_and_html_is_escaped(self):
        employee = '=HYPERLINK("https://example.invalid")'
        self.add_shift("2026-10-07", employee, "work", 1)
        self.add_shift("2026-10-07", "<script>alert(1)</script>", "work", 1)
        response = self.client.get("/object-report.xlsx", query_string=self.params())
        workbook = load_workbook(io.BytesIO(response.data), data_only=False)
        for sheet in workbook:
            self.assertTrue(all(cell.data_type != "f" for row in sheet for cell in row))
        response = self.client.get("/object-report", query_string=self.params())
        self.assertIn("&lt;script&gt;alert(1)&lt;/script&gt;", response.get_data(as_text=True))

    def test_invalid_server_data_is_reported_not_silently_skipped(self):
        for kind, hours in (("unknown", 8), ("work", -1), ("work", "inf"), ("work", "bad")):
            self.add_shift("2026-10-10", "Anna Åström", kind, hours)
            with self.assertLogs("object_reports", "ERROR"):
                response = self.client.get("/object-report.xlsx", query_string=self.params())
            self.assertEqual(response.status_code, 500)
            self.assertIn("ошибка записана", response.get_json()["error"])
            with closing(self.connection()) as conn:
                conn.execute("DELETE FROM incoming_submissions WHERE date='2026-10-10'")
                conn.commit()

    def test_empty_report_and_multi_page_pdf(self):
        for suffix in ("", ".pdf", ".xlsx"):
            response = self.client.get(
                "/object-report" + suffix, query_string=self.params(start="2027-01-01", end="2027-01-31")
            )
            self.assertEqual(response.status_code, 200)
        for index in range(110):
            self.add_shift("2026-10-10", f"Medarbetare {index:03}", "work", 0.5)
        response = self.client.get("/object-report.pdf", query_string=self.params())
        reader = PdfReader(io.BytesIO(response.data))
        self.assertGreater(len(reader.pages), 2)
        self.assertIn("Medarbetare 109", "\n".join(p.extract_text() for p in reader.pages))

    def test_export_failure_is_logged_and_explicit(self):
        with patch("object_reports.report_pdf", side_effect=OSError("isolated export failure")), \
                self.assertLogs("object_reports", "ERROR"):
            response = self.client.get("/object-report.pdf", query_string=self.params())
        self.assertEqual(response.status_code, 500)
        self.assertIn("Не удалось сформировать", response.get_json()["error"])

    def test_missing_database_is_not_created_and_post_is_not_allowed(self):
        self.database = Path(self.directory.name) / "does-not-exist.db"
        with self.assertLogs("object_reports", "ERROR"):
            self.assertEqual(self.client.get("/object-report").status_code, 500)
        self.assertFalse(self.database.exists())
        self.assertEqual(self.client.post("/object-report").status_code, 405)

    def test_full_server_integration_on_disposable_database(self):
        env = {**os.environ, "DATABASE_PATH": str(Path(self.directory.name) / "integration.db"),
               "ADMIN_PASSWORD": "isolated-report-only", "SECRET_KEY": "isolated-report-secret",
               "SYNC_API_KEY": "isolated-report-sync"}
        result = subprocess.run([sys.executable, "-c", """
import server
client = server.app.test_client()
assert client.get('/object-report').status_code == 401
response = client.post('/login', json={'username':'admin','password':'isolated-report-only'})
assert response.status_code == 200, response.get_json()
assert client.get('/object-report').status_code == 200
assert b'object-report-link' in client.get('/').data
connection = server.get_db_connection()
connection.execute("CREATE TABLE sync_dataset(id INTEGER PRIMARY KEY, dataset_id TEXT)")
connection.execute("INSERT INTO sync_dataset VALUES (1, 'isolated-report-dataset')")
connection.commit()
assert client.get('/object-report').status_code == 401
assert client.post('/login', json={'username':'admin','password':'isolated-report-only'}).status_code == 200
assert client.get('/object-report').status_code == 200
connection.execute("INSERT INTO meta_employees(name) VALUES ('admin')")
connection.execute("UPDATE users SET role='user' WHERE username='admin'")
connection.commit()
assert client.get('/object-report').status_code == 403
connection.execute("UPDATE users SET is_active=0 WHERE username='admin'")
connection.commit()
connection.close()
assert client.get('/object-report').status_code == 401
"""], cwd=ROOT, env=env, capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr[-4000:])


if __name__ == "__main__":
    unittest.main()
