import calendar
import hashlib
import io
import json
import logging
import os
import re
import sqlite3
from collections import defaultdict
from contextlib import closing
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Callable
from xml.sax.saxutils import escape

import reportlab
from flask import Blueprint, Flask, jsonify, render_template, request, send_file, session, url_for
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle
from werkzeug.utils import secure_filename


logger = logging.getLogger(__name__)
TYPE_LABELS = {
    "work": "Arbete",
    "ob": "OB / Övertid",
    "atf": "ATF",
    "sick": "Sjuk",
    "vacation": "Semester",
    "holiday": "Röda dagar",
}
TYPE_LABELS_RU = {
    "work": "Работа",
    "ob": "OB / сверхурочные",
    "atf": "АТФ",
    "sick": "Больничные",
    "vacation": "Отпуск",
    "holiday": "Красные дни",
}
SOURCE_NOTE = (
    "Rapporten omfattar endast serverdata. Poster som ännu inte skickats från "
    "arbetsdatorn ingår inte. Uppgifterna kan ändras mellan hämtningar."
)
ROUNDING_NOTE = (
    "Timmar avrundas till 0,01 per datum, medarbetare och typ. "
    "Summorna avser de visade posterna. OB-koefficienter tillämpas inte på timmar."
)
ACCOUNTING_NOTE_RU = (
    "Часы приведены по записям учёта. Пересечения обычной работы с OB автоматически "
    "не проверяются, поскольку для обычной работы отсутствуют время начала и окончания"
)
ACCOUNTING_NOTE = (
    "Timmarna anges enligt registrerade poster. Överlappningar mellan vanligt arbete "
    "och OB kontrolleras inte automatiskt eftersom vanligt arbete saknar start- och sluttid."
)
OVERLAP_RU = (
    "Предупреждение: обнаружены пересекающиеся OB-интервалы одного сотрудника, "
    "затрагивающие записи выбранного объекта. Проверка выполнена по всем объектам. "
    "Часы не объединены и не вычтены; итог является суммой записей."
)
OVERLAP_SV = (
    "Varning: överlappande OB-intervall för samma medarbetare berör poster i det valda "
    "objektet. Kontrollen omfattar alla objekt. Inga timmar har slagits samman eller "
    "dragits av; summan avser registrerade poster."
)
INCOMPLETE_RU = (
    "Предупреждение: проверка пересечений OB неполная — у части записей отсутствует "
    "или некорректно время начала/окончания. Часы сохранены в отчёте без изменений."
)
INCOMPLETE_SV = (
    "Varning: kontrollen av OB-överlappningar är ofullständig eftersom vissa poster "
    "saknar giltig start- eller sluttid. Timmarna i rapporten har inte ändrats."
)
HISTORY_NOTE = (
    "Объект определяется точной парой «заказчик + название на сервере». "
    "Старые названия без подтверждённой связи показаны отдельно как «только история» "
    "и не объединяются автоматически."
)
ZERO = Decimal("0.00")


class ReportInputError(ValueError):
    pass


class ReportDataError(ValueError):
    pass


@dataclass(frozen=True)
class ReportObject:
    key: str
    name: str
    company: str
    marking: str
    state: str

    @property
    def label(self) -> str:
        return f"{self.marking + ' | ' if self.marking else ''}{self.name} [{self.company}]"


@dataclass(frozen=True)
class ReportRow:
    day: str
    employee: str
    kind: str
    hours: Decimal


@dataclass
class EmployeeTotal:
    employee: str
    hours: dict[str, Decimal]

    @property
    def total(self) -> Decimal:
        return sum(self.hours.values(), ZERO)


@dataclass
class ObjectReport:
    project: ReportObject
    start: date
    end: date
    generated_at: str
    rows: list[ReportRow]
    employees: list[EmployeeTotal]
    totals: EmployeeTotal
    ob_overlap: bool
    ob_check_incomplete: bool

    @property
    def warnings_ru(self) -> list[str]:
        return ([OVERLAP_RU] if self.ob_overlap else []) + (
            [INCOMPLETE_RU] if self.ob_check_incomplete else []
        )

    @property
    def warnings_sv(self) -> list[str]:
        return ([OVERLAP_SV] if self.ob_overlap else []) + (
            [INCOMPLETE_SV] if self.ob_check_incomplete else []
        )


def _ob_interval(day: str, start: str, end: str) -> tuple[datetime, datetime]:
    if not re.fullmatch(r"(?:[01]\d|2[0-3]):[0-5]\d", start or ""):
        raise ValueError("Invalid OB start")
    if not re.fullmatch(r"(?:(?:[01]\d|2[0-3]):[0-5]\d|24:00)", end or "") or start == end:
        raise ValueError("Invalid OB end")
    midnight = datetime.combine(date.fromisoformat(day), datetime.min.time())
    start_minutes = int(start[:2]) * 60 + int(start[3:])
    end_minutes = int(end[:2]) * 60 + int(end[3:])
    if end_minutes < start_minutes:
        end_minutes += 24 * 60
    return midnight + timedelta(minutes=start_minutes), midnight + timedelta(minutes=end_minutes)


def check_ob_overlap(
    conn: sqlite3.Connection, project: ReportObject, start: date, end: date,
) -> tuple[bool, bool]:
    normalize = lambda name: " ".join(str(name or "").split()).casefold()
    conn.create_function("report_employee_key", 1, normalize, deterministic=True)
    lower = start - timedelta(days=1) if start > date.min else start
    upper = end + timedelta(days=1) if end < date.max else end
    records = conn.execute("""
        SELECT id, employee, date, start_time, end_time, company, object_name
        FROM incoming_submissions AS candidate
        WHERE submission_type = 'ob' AND is_deleted = 0 AND date >= ? AND date <= ?
          AND EXISTS (
              SELECT 1 FROM incoming_submissions AS selected
              WHERE selected.submission_type = 'ob' AND selected.is_deleted = 0
                AND selected.company = ? AND selected.object_name = ?
                AND selected.date >= ? AND selected.date <= ?
                AND report_employee_key(selected.employee) = report_employee_key(candidate.employee)
          )
    """, (lower.isoformat(), upper.isoformat(), project.company, project.name, start.isoformat(), end.isoformat()))
    by_employee: dict[str, list[tuple[datetime, datetime, bool]]] = defaultdict(list)
    incomplete = False
    for record in records:
        selected = (
            record["company"] == project.company and record["object_name"] == project.name
            and start.isoformat() <= record["date"] <= end.isoformat()
        )
        try:
            begin, finish = _ob_interval(record["date"], record["start_time"], record["end_time"])
        except (ValueError, OverflowError):
            incomplete = True
            continue
        by_employee[normalize(record["employee"])].append((begin, finish, selected))
    for intervals in by_employee.values():
        latest_end = latest_selected_end = datetime.min
        for begin, finish, selected in sorted(intervals):
            if begin < latest_selected_end or (selected and begin < latest_end):
                return True, incomplete
            latest_end = max(latest_end, finish)
            if selected:
                latest_selected_end = max(latest_selected_end, finish)
    return False, incomplete


def report_period(preset: str, start_text: str, end_text: str) -> tuple[date, date]:
    try:
        start = date.fromisoformat(start_text)
        if start.isoformat() != start_text:
            raise ValueError()
        if preset == "custom":
            end = date.fromisoformat(end_text)
            if end.isoformat() != end_text:
                raise ValueError()
        elif preset in {"week", "fortnight"}:
            start -= timedelta(days=start.weekday())
            end = start + timedelta(days=6 if preset == "week" else 13)
        elif preset == "month":
            start = start.replace(day=1)
            end = start.replace(day=calendar.monthrange(start.year, start.month)[1])
        else:
            raise ReportInputError("Неизвестный период.")
    except (ValueError, OverflowError) as error:
        raise ReportInputError("Укажите допустимые даты в формате ГГГГ-ММ-ДД и период.") from error
    if end < start:
        raise ReportInputError("Конец периода не может быть раньше начала.")
    return start, end


def _object_key(company: str, name: str) -> str:
    return hashlib.sha256(json.dumps([company, name], ensure_ascii=False).encode("utf-8")).hexdigest()


def report_objects(conn: sqlite3.Connection) -> list[ReportObject]:
    projects = {}
    for row in conn.execute(
        "SELECT name, COALESCE(company, '') AS company, COALESCE(markning, '') AS marking, "
        "is_active FROM meta_objects ORDER BY id"
    ):
        company, name = row["company"], row["name"]
        projects[(company, name)] = ReportObject(
            _object_key(company, name), name, company, row["marking"],
            "" if row["is_active"] else "неактивен",
        )
    for row in conn.execute(
        "SELECT DISTINCT company, object_name FROM incoming_submissions WHERE is_deleted = 0"
    ):
        company, name = row["company"], row["object_name"]
        if (company, name) not in projects:
            projects[(company, name)] = ReportObject(
                _object_key(company, name), name, company, "", "только история",
            )
    return sorted(projects.values(), key=lambda p: (p.company.casefold(), p.name.casefold()))


def build_report(conn: sqlite3.Connection, project: ReportObject, start: date, end: date) -> ObjectReport:
    grouped: dict[tuple[str, str, str], Decimal] = defaultdict(lambda: ZERO)
    names = {}
    # One canonical table: cloud_shifts is a legacy mirror, not additional hours.
    records = conn.execute("""
        SELECT date, employee, submission_type, hours
        FROM incoming_submissions
        WHERE company = ? AND object_name = ? AND date >= ? AND date <= ? AND is_deleted = 0
        ORDER BY date, employee, id
    """, (project.company, project.name, start.isoformat(), end.isoformat()))
    for record in records:
        kind = record["submission_type"]
        employee = " ".join(record["employee"].split())
        if kind not in TYPE_LABELS or not employee:
            raise ReportDataError("В серверных записях обнаружен неизвестный тип часов или пустой сотрудник.")
        try:
            hours = Decimal(str(record["hours"]))
        except InvalidOperation as error:
            raise ReportDataError("В серверных записях обнаружено некорректное количество часов.") from error
        if not hours.is_finite() or hours < 0:
            raise ReportDataError("В серверных записях обнаружено некорректное количество часов.")
        key = employee.casefold()
        names.setdefault(key, employee)
        grouped[(record["date"], key, kind)] += hours

    rows = [
        ReportRow(day, names[employee], kind, hours.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))
        for (day, employee, kind), hours in sorted(grouped.items())
    ]
    by_employee: dict[str, dict[str, Decimal]] = {}
    for row in rows:
        totals = by_employee.setdefault(row.employee, dict.fromkeys(TYPE_LABELS, ZERO))
        totals[row.kind] += row.hours
    employees = [
        EmployeeTotal(name, hours)
        for name, hours in sorted(by_employee.items(), key=lambda item: item[0].casefold())
    ]
    totals = EmployeeTotal("Totalt", {
        kind: sum((employee.hours[kind] for employee in employees), ZERO) for kind in TYPE_LABELS
    })
    overlap, incomplete = check_ob_overlap(conn, project, start, end)
    return ObjectReport(
        project, start, end, datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC"),
        rows, employees, totals, overlap, incomplete,
    )


def _pdf_font() -> str:
    name = "ObjectReport"
    if name not in pdfmetrics.getRegisteredFontNames():
        candidates = [
            os.environ.get("ARBETSDAGBOK_FONT_PATH"),
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            str(Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts" / "arial.ttf"),
            str(Path(reportlab.__file__).parent / "fonts" / "Vera.ttf"),
        ]
        font = next((p for p in candidates if p and Path(p).is_file()), None)
        if font is None:
            raise RuntimeError("Не найден шрифт для PDF-отчёта.")
        pdfmetrics.registerFont(TTFont(name, font))
    return name


def report_pdf(report: ObjectReport) -> io.BytesIO:
    font = _pdf_font()
    style = ParagraphStyle("ObjectReportCell", fontName=font, fontSize=8, leading=11)
    heading = ParagraphStyle(
        "ObjectReportHeading", parent=getSampleStyleSheet()["Heading1"], fontName=font,
        fontSize=15, leading=19,
    )
    warning_style = ParagraphStyle(
        "ObjectReportWarning", parent=style, backColor=colors.HexColor("#fff3cd"),
        textColor=colors.HexColor("#694c00"), borderPadding=6, spaceBefore=6, spaceAfter=6,
    )

    def paragraph(text: object) -> Paragraph:
        return Paragraph(escape(str(text)), style)

    def table(data: list[list[object]], widths: list[float]) -> LongTable:
        result = LongTable([[paragraph(cell) for cell in row] for row in data], colWidths=widths, repeatRows=1)
        result.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8edf2")),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#b8c2cc")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        return result

    story = [
        Paragraph("Objektrapport / Timmar", heading),
        paragraph(report.project.label),
        paragraph(f"Period: {report.start.isoformat()} – {report.end.isoformat()} (inklusive)"),
        paragraph(f"Skapad: {report.generated_at}"),
        paragraph(SOURCE_NOTE),
        paragraph(ROUNDING_NOTE),
        paragraph(ACCOUNTING_NOTE),
        Spacer(1, 4 * mm),
        paragraph(f"Totalt registrerade timmar: {report.totals.total:.2f} h"),
        Spacer(1, 4 * mm),
    ]
    for warning in report.warnings_sv:
        story.extend([Paragraph(escape(warning), warning_style), Spacer(1, 3 * mm)])
    summary: list[list[object]] = [[
        "Medarbetare", *TYPE_LABELS.values(), "Totalt registrerade timmar",
    ]]
    for employee in [*report.employees, report.totals]:
        summary.append([
            employee.employee, *[f"{employee.hours[k]:.2f}" for k in TYPE_LABELS],
            f"{employee.total:.2f}",
        ])
    story.extend([table(summary, [52 * mm] + [28 * mm] * 6 + [32 * mm]), Spacer(1, 6 * mm)])
    details: list[list[object]] = [["Datum", "Medarbetare", "Typ", "Timmar"]]
    details.extend([[r.day, r.employee, TYPE_LABELS[r.kind], f"{r.hours:.2f}"] for r in report.rows])
    if not report.rows:
        story.append(paragraph("Inga poster för vald period."))
    else:
        story.append(table(details, [28 * mm, 124 * mm, 75 * mm, 25 * mm]))
    output = io.BytesIO()
    doc = SimpleDocTemplate(
        output, pagesize=landscape(A4), leftMargin=18 * mm, rightMargin=18 * mm,
        topMargin=15 * mm, bottomMargin=15 * mm, title="Objektrapport", author="SmartReport",
    )

    def page_number(canvas, document):
        canvas.setFont(font, 8)
        canvas.drawRightString(279 * mm, 8 * mm, f"SmartReport | {document.page}")

    doc.build(story, onFirstPage=page_number, onLaterPages=page_number)
    output.seek(0)
    return output


def report_xlsx(report: ObjectReport) -> io.BytesIO:
    workbook = Workbook()
    summary = workbook.active
    if summary is None:
        raise RuntimeError("Не удалось создать лист Excel.")
    summary.title = "Sammanfattning"
    for line in [
        ["Objektrapport / Timmar"], [report.project.label],
        [f"Period: {report.start.isoformat()} – {report.end.isoformat()} (inklusive)"],
        [f"Skapad: {report.generated_at}"], [SOURCE_NOTE], [ROUNDING_NOTE],
        [ACCOUNTING_NOTE], *[[warning] for warning in report.warnings_sv], [],
        ["Medarbetare", *TYPE_LABELS.values(), "Totalt registrerade timmar"],
    ]:
        summary.append(line)
    summary_header = summary.max_row
    for employee in [*report.employees, report.totals]:
        summary.append([
            employee.employee, *[employee.hours[k] for k in TYPE_LABELS],
            employee.total,
        ])
    details = workbook.create_sheet("Timmar")
    details.append(["Datum", "Medarbetare", "Typ", "Timmar"])
    for row in report.rows:
        details.append([date.fromisoformat(row.day), row.employee, TYPE_LABELS[row.kind], row.hours])
    for sheet, header, widths in [
        (summary, summary_header, [42, 15, 18, 12, 12, 15, 17, 28]),
        (details, 1, [16, 42, 24, 16]),
    ]:
        sheet.freeze_panes = f"B{header + 1}"
        sheet.auto_filter.ref = f"A{header}:{sheet.cell(sheet.max_row, sheet.max_column).coordinate}"
        for column, width in enumerate(widths, 1):
            sheet.column_dimensions[sheet.cell(header, column).column_letter].width = width
        for row in sheet:
            for cell in row:
                # Export user-supplied names literally, never as spreadsheet formulas.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                elif isinstance(cell.value, (Decimal, float, int)):
                    cell.number_format = "0.00"
                elif isinstance(cell.value, date):
                    cell.number_format = "yyyy-mm-dd"
                if cell.row == header:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="24527A")
        if sheet == summary:
            for row in range(2, summary_header - 1):
                sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
                sheet.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="top")
                sheet.row_dimensions[row].height = 32
                if row >= 8:
                    sheet.row_dimensions[row].height = 44
                    sheet.cell(row, 1).fill = PatternFill("solid", fgColor="FFF3CD")
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def register_object_reports(app: Flask, database_path: Callable[[], str]) -> None:
    blueprint = Blueprint("object_reports", __name__)

    @blueprint.after_request
    def no_report_cache(response):
        response.headers["Cache-Control"] = "no-store, private"
        return response

    @blueprint.route("/object-report", defaults={"output_format": "html"})
    @blueprint.route("/object-report.<output_format>")
    def object_report(output_format):
        if not session.get("user_id"):
            return jsonify(ok=False, error="Требуется вход администратора."), 401
        if session.get("role") != "admin":
            return jsonify(ok=False, error="Отчёт доступен только администратору."), 403
        if output_format not in {"html", "pdf", "xlsx"}:
            return jsonify(ok=False, error="Неизвестный формат отчёта."), 404
        values = {
            "project": request.args.get("project", ""),
            "preset": request.args.get("preset", "month"),
            "start": request.args.get("start", datetime.now(timezone.utc).date().isoformat()),
            "end": request.args.get("end", ""),
        }
        projects = []
        report = None
        error = None
        status = 200
        try:
            uri = Path(database_path()).resolve().as_uri() + "?mode=ro"
            with closing(sqlite3.connect(uri, uri=True, timeout=30)) as conn:
                conn.row_factory = sqlite3.Row
                conn.execute("PRAGMA query_only=ON")
                conn.execute("BEGIN")
                projects = report_objects(conn)
                if values["project"]:
                    project = next((p for p in projects if p.key == values["project"]), None)
                    if project is None:
                        raise ReportInputError("Объект не найден. Выберите объект из списка.")
                    start, end = report_period(values["preset"], values["start"], values["end"])
                    report = build_report(conn, project, start, end)
                elif output_format != "html" or request.args:
                    raise ReportInputError("Выберите объект.")
        except ReportInputError as exc:
            error, status = str(exc), 400
        except (sqlite3.Error, ReportDataError, InvalidOperation):
            logger.exception("Не удалось прочитать данные отчёта по объекту")
            error, status = "Не удалось прочитать отчёт. Проверьте серверные данные; ошибка записана в журнал.", 500
        if error and output_format != "html":
            return jsonify(ok=False, error=error), status
        if output_format == "html":
            export_args = {
                **values, "preset": "custom", "start": report.start.isoformat(), "end": report.end.isoformat(),
            } if report else {}
            return render_template(
                "object_report.html", projects=projects, report=report, values=values, error=error,
                type_labels=TYPE_LABELS_RU, history_note=HISTORY_NOTE, accounting_note=ACCOUNTING_NOTE_RU,
                pdf_url=url_for(".object_report", output_format="pdf", **export_args) if report else "",
                xlsx_url=url_for(".object_report", output_format="xlsx", **export_args) if report else "",
            ), status
        if report is None:
            logger.error("Файл отчёта запрошен без сформированного отчёта")
            return jsonify(ok=False, error="Отчёт не сформирован."), 500
        try:
            output = report_pdf(report) if output_format == "pdf" else report_xlsx(report)
        except Exception:
            logger.exception("Не удалось сформировать отчёт по объекту в формате %s", output_format)
            return jsonify(ok=False, error="Не удалось сформировать файл отчёта; ошибка записана в журнал."), 500
        mimetype = (
            "application/pdf" if output_format == "pdf"
            else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        project_name = secure_filename((report.project.marking or report.project.name)[:60])
        project_name = project_name or report.project.key[:12]
        return send_file(
            output, mimetype=mimetype, as_attachment=True,
            download_name=(
                f"Objektrapport_{project_name}_{report.start.isoformat()}_"
                f"{report.end.isoformat()}.{output_format}"
            ),
        )

    app.register_blueprint(blueprint)
